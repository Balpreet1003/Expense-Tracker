from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.features.auth.models.user import User
from app.features.income.models.income import Income


def download_incomes(
    db: Session,
    current_user: User,
) -> BytesIO:

    statement = (
        select(Income)
        .where(
            Income.user_id == current_user.id
        )
        .order_by(
            Income.date.desc(),
        )
    )

    result = db.execute(statement)

    incomes = result.scalars().all()

    workbook = Workbook()

    worksheet = workbook.active

    worksheet.title = "Incomes"

    headers = [
        "ID",
        "Icon",
        "Amount",
        "Date",
        "Source",
        "Description",
    ]

    worksheet.append(headers)

    # Make headers bold
    for cell in worksheet[1]:
        cell.font = Font(bold=True)

    # Add income data
    for income in incomes:
        worksheet.append(
            [
                income.id,
                income.icon,
                income.amount,
                income.date,
                income.source,
                income.description,
            ]
        )

    # Adjust column widths
    for column_cells in worksheet.columns:
        max_length = max(
            len(str(cell.value))
            if cell.value is not None
            else 0
            for cell in column_cells
        )

        column_letter = get_column_letter(
            column_cells[0].column
        )

        worksheet.column_dimensions[
            column_letter
        ].width = max_length + 2

    output = BytesIO()

    workbook.save(output)

    output.seek(0)

    return output
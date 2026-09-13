from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.features.auth.models.user import User
from app.features.expense.models.expense import Expense


def download_expenses(
    db: Session,
    current_user: User,
) -> BytesIO:

    statement = (
        select(Expense)
        .where(
            Expense.user_id == current_user.id
        )
        .order_by(
            Expense.date.desc(),
        )
    )

    result = db.execute(statement)

    expenses = result.scalars().all()

    workbook = Workbook()

    worksheet = workbook.active

    worksheet.title = "Expenses"

    headers = [
        "ID",
        "Icon",
        "Amount",
        "Date",
        "Category",
        "Description",
    ]

    worksheet.append(headers)

    # Make headers bold
    for cell in worksheet[1]:
        cell.font = Font(bold=True)

    # Add expense data
    for expense in expenses:
        worksheet.append(
            [
                expense.id,
                expense.icon,
                expense.amount,
                expense.date,
                expense.category,
                expense.description,
            ]
        )

    # Adjust column widths
    for column_cells in worksheet.columns:
        max_length = max(
            len(str(cell.value))
            if cell.value is not None
            else 0
            for cell in column_cells
        )

        column_letter = get_column_letter(
            column_cells[0].column
        )

        worksheet.column_dimensions[
            column_letter
        ].width = max_length + 2

    output = BytesIO()

    workbook.save(output)

    output.seek(0)

    return output
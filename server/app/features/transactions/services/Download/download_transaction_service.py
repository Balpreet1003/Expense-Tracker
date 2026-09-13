from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from app.features.auth.models.user import User
from app.features.transactions.services.Get.get_transaction_service import (
    get_transactions,
)


def download_transactions(
    db: Session,
    current_user: User,
) -> BytesIO:

    transactions = get_transactions(
        db=db,
        current_user=current_user,
    )

    workbook = Workbook()

    worksheet = workbook.active

    worksheet.title = "Transactions"

    headers = [
        "ID",
        "Type",
        "Icon",
        "Amount",
        "Date",
        "Source/Category",
        "Description",
    ]

    worksheet.append(headers)

    # Make headers bold
    for cell in worksheet[1]:
        cell.font = Font(bold=True)

    for transaction in transactions:

        source_or_category = (
            transaction.source
            if transaction.type == "income"
            else transaction.category
        )

        worksheet.append(
            [
                transaction.id,
                transaction.type,
                transaction.icon,
                transaction.amount,
                transaction.date,
                source_or_category,
                transaction.description,
            ]
        )

    # Adjust column widths
    for column_cells in worksheet.columns:
        max_length = max(
            len(str(cell.value))
            if cell.value is not None
            else 0
            for cell in column_cells
        )

        column_letter = get_column_letter(
            column_cells[0].column
        )

        worksheet.column_dimensions[
            column_letter
        ].width = max_length + 2

    output = BytesIO()

    workbook.save(output)

    output.seek(0)

    return output